"""Parser of the organizers' incident classifier (xlsx, sheet «Лист1»).

Layout (see docs/DATASET.md):

* rows 1-3 are a three-level header. Columns A-D hold the codes «Г, п1, п2, п3», F the
  statistics group, G-I the three levels of formalized signs («112-Признак.1/2/3»), J free-text
  hints, K the final incident type, L the ЕКП type, M the main service;
* from column N on every column belongs to a service. Row 1 names the service family,
  row 2 a sub-service, row 3 the flag condition («выбран признак НД», «признак не выбран»…);
* data starts at row 4. A row with an empty K and a filled F opens an incident group;
  every other non-empty row is one final incident type.

A service is on the notification list of a type when its cell is filled. A cell equal to
«нет реагирования» is the opposite: the service is explicitly *not* notified under that
flag. Flag columns add services when the flag is chosen; default columns («признак не
выбран») apply always. See ``app.domain.services.resolve_services``.
"""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

SHEET_NAME = "Лист1"
HEADER_ROWS = 3
FIRST_SERVICE_COLUMN = 14  # N
COL_G, COL_P1, COL_P2, COL_P3, COL_NUMBER = 1, 2, 3, 4, 5
COL_STAT_GROUP, COL_SIGN1, COL_SIGN2, COL_SIGN3, COL_HINTS = 6, 7, 8, 9, 10
COL_FINAL, COL_EKP, COL_MAIN = 11, 12, 13

NO_RESPONSE = "нет реагирования"

# (family header, sub header) → (service code, title, short title).
# Headers are matched after whitespace normalization and lower-casing.
SERVICE_COLUMNS: dict[tuple[str, str], tuple[str, str, str]] = {
    ("классификатор мчс", "служба 101"): ("101", "Служба 101 (МЧС)", "101"),
    ("классификатор мчс", "одс псц"): ("ods_psc", "ОДС ПСЦ (Пожарно-спасательный центр)", "ПСЦ"),
    ("классификатор мчс", "мгпсс"): ("mgpss", "МГПСС (спасение на воде)", "МГПСС"),
    ("классификатор мвд", ""): ("102", "Служба 102 (МВД)", "102"),
    ("классификатор смп", ""): ("103", "Служба 103 (скорая помощь)", "103"),
    ("классификатор мосгаз", ""): ("mosgaz", "Мосгаз", "Мосгаз"),
    ("цэмп", ""): ("cemp", "ЦЭМП (Центр экстренной медицинской помощи)", "ЦЭМП"),
    ("классификатор фсб", ""): ("fsb", "ФСБ", "ФСБ"),
    ("классификатор мособлгаз", ""): ("mosoblgaz", "Мособлгаз", "Мособлгаз"),
    ("автомобильные дороги", ""): ("autoroads", "ГБУ «Автомобильные дороги»", "Автодороги"),
    ("мосгортранс", ""): ("mosgortrans", "Мосгортранс", "Мосгортранс"),
    ("гор. хозяйство", ""): ("gkh", "Городское хозяйство (ЕДЦ ЖКХ)", "Гор. хозяйство"),
    ("гормост", ""): ("gormost", "ГБУ «Гормост»", "Гормост"),
    ("канал имени москвы", ""): ("kim", "Канал имени Москвы", "КиМ"),
    ("мгтс", ""): ("mgts", "МГТС", "МГТС"),
    ("метро", ""): ("metro", "Московский метрополитен", "Метро"),
    ("мосводоканал", ""): ("mosvodokanal", "Мосводоканал", "Мосводоканал"),
    ("моэк", ""): ("moek", "МОЭК", "МОЭК"),
    ('моэск (пао "россети московский регион")', ""): (
        "moesk",
        "Россети Московский регион (МОЭСК)",
        "МОЭСК",
    ),
    ("оэк", ""): ("oek", "ОЭК", "ОЭК"),
    ("мослифт", ""): ("moslift", "Мослифт", "Мослифт"),
    ("цодд", ""): ("codd", "ЦОДД", "ЦОДД"),
    ("деп. жкх", ""): ("dep_gkh", "Департамент ЖКХ", "Деп. ЖКХ"),
    ("департамент рбипк (гку мосбез)", "дежурная служба арм-112"): (
        "duty_arm112",
        "Дежурная служба АРМ-112 (ГКУ МОСБЕЗ)",
        "ДС АРМ-112",
    ),
    ("департамент рбипк (гку мосбез)", "мкп, аналитика (старый крим)"): (
        "mkp_analytics",
        "МКП, Аналитика (ГКУ МОСБЕЗ)",
        "МКП",
    ),
    ("аппарат мэра", ""): ("mayor_office", "Аппарат Мэра", "Аппарат Мэра"),
    ("москоллектор", ""): ("moscollector", "Москоллектор", "Москоллектор"),
    ("ржд (московскаяжд)", ""): ("mzd", "РЖД (Московская железная дорога)", "МЖД"),
    ("департамент образования", ""): ("dep_education", "Департамент образования", "Деп. обр."),
    ("центррегионводхоз (московско-окское бву)", ""): (
        "crvh",
        "Центррегионводхоз (Московско-Окское БВУ)",
        "ЦРВХ",
    ),
    ("военная комендатура", ""): ("military_commandant", "Военная комендатура", "Комендатура"),
    ("оати", ""): ("oati", "ОАТИ", "ОАТИ"),
    ("мосводосток", ""): ("mosvodostok", "Мосводосток", "Мосводосток"),
    ("департамент ппиоос", ""): (
        "dep_ppoos",
        "Департамент природопользования и охраны окружающей среды",
        "Деп. природ.",
    ),
    ("од департамент тсзн", ""): (
        "dep_tszn",
        "Департамент труда и социальной защиты населения",
        "Деп. ТСЗН",
    ),
    ("рсво", ""): ("rsvo", "РСВО (Российские сети вещания и оповещения)", "РСВО"),
    ("эважд", ""): ("evazhd", "ЭВАЖД (эксплуатация высотных домов)", "ЭВАЖД"),
    ("мсппн", ""): ("msppn", "МСППН (психологическая помощь населению)", "МСППН"),
    ("дту_р (ритуал)", ""): ("dtu_ritual", "ДТУ (Ритуал)", "Ритуал"),
    ("дту", ""): ("dtu", "ДТУ (Департамент торговли и услуг)", "ДТУ"),
    ("росгвардия", ""): ("rosgvardia", "Росгвардия", "Росгвардия"),
    ("территориальные оив", ""): (
        "territorial_oiv",
        "Территориальные ОИВ (управы, префектуры)",
        "Управа",
    ),
    ("территориальные оив тинао", ""): (
        "territorial_oiv_tinao",
        "Территориальные ОИВ ТиНАО",
        "Управа ТиНАО",
    ),
    ("автомобильные дороги ао г.москвы", ""): (
        "autoroads_ao",
        "Автомобильные дороги АО г. Москвы",
        "Автодороги АО",
    ),
    ("департамент строительства города москвы", ""): (
        "dep_construction",
        "Департамент строительства города Москвы",
        "Деп. строит.",
    ),
    ("комитет ветеринарии", ""): ("veterinary", "Комитет ветеринарии", "Ветеринария"),
    ("мосжилинспекция", ""): ("moszhilinspection", "Мосжилинспекция", "МЖИ"),
    ("департамент культуры", ""): ("dep_culture", "Департамент культуры", "Деп. культуры"),
    ("гку цса имени е.п.глинки", ""): ("csa_glinki", "ГКУ ЦСА имени Е. П. Глинки", "ЦСА"),
    ("гку нту", ""): ("ntu", "ГКУ НТУ", "НТУ"),
    ("фсо", ""): ("fso", "ФСО", "ФСО"),
    ("гуп мср", "(куб)"): ("msr_kub", "ГУП МСР (КУБ)", "МСР КУБ"),
    ("гуп мср", "пожары"): ("msr_fires", "ГУП МСР (пожары)", "МСР пожары"),
    ("комитет по туризму г.москвы", ""): ("tourism", "Комитет по туризму г. Москвы", "Туризм"),
    ("дгп (департамент градостроительной политики)", "интеграция"): (
        "dgp_integration",
        "ДГП (интеграция)",
        "ДГП",
    ),
    ("дгп (департамент градостроительной политики)", "арм-112"): (
        "dgp_arm112",
        "ДГП (АРМ-112)",
        "ДГП АРМ-112",
    ),
    ("цукб министерство обороны", ""): ("cukb_mo", "ЦУКБ Министерства обороны", "ЦУКБ МО"),
    ("цукб.бпла министерство обороны", ""): (
        "cukb_bpla_mo",
        "ЦУКБ.БПЛА Министерства обороны",
        "ЦУКБ БПЛА",
    ),
    ("гку организатор перевозок", ""): (
        "transport_organizer",
        "ГКУ «Организатор перевозок»",
        "Орг. перевозок",
    ),
    ("гпбу мосэкомониторинг", ""): (
        "mosecomonitoring",
        "ГПБУ «Мосэкомониторинг»",
        "Мосэкомониторинг",
    ),
    ("министерство обороны рхбз", "события по полигонам"): (
        "mo_rhbz_polygons",
        "Минобороны РХБЗ (события по полигонам)",
        "РХБЗ полигоны",
    ),
    ("министерство обороны рхбз", "москва"): (
        "mo_rhbz_moscow",
        "Минобороны РХБЗ (Москва)",
        "РХБЗ Москва",
    ),
    ("ооо ситиэнерго", ""): ("cityenergo", "ООО «Ситиэнерго»", "Ситиэнерго"),
    ("депортамент гражданского строительства", ""): (
        "dep_civil_construction",
        "Департамент гражданского строительства",
        "Деп. гражд. стр.",
    ),
}

# Row-3 condition text → flag code. None means the column applies regardless of flags.
CONDITION_FLAGS: dict[str, str | None] = {
    "": None,
    "признак не выбран": None,
    "признаки не выбраны": None,
    "реагирование всегда": None,
    "служба 101 (признак нд - нет доступа не выбран)": None,
    "служба 101 (выбран признак нд - нет доступа)": "no_access",
    "одс псц (другие признаки не выбраны)": None,
    "одс псц (выбран признак ул - угроза людям)": "threat",
    "одс псц (выбран признак пп - пострадавшие погибшие)": "injured",
    "одс псц (выбран признак нд - нет доступа)": "no_access",
    "признак правонарушение или пострадавшие не выбран": None,
    "выбран признак правонарушение": "offense",
    "выбран признак пострадавшие": "injured",
    "классификатор смп (признак пострадавшие не выбран)": None,
    "классификатор смп (выбран признак пострадавшие)": "injured",
    "классификатор смп (выбран признак пострадавшие не на месте)": "not_on_site",
    "газификация": "gasification",
    "угроза людям": "threat",
    "пострадавшие/погибшие": "injured",
    "мед. помощь": "medical_help",
    "треб. эвакуация": "evacuation",
    ">5 чел / од": "mass_incident",
    "постр / погибшие": "injured",
    "перекрытие движение": "road_closed",
    "тоннель": "tunnel",
    "пеш": "pedestrian_bridge",
    "ав": "road_bridge",
    "на объектах связи": "telecom_object",
    "стройка": "construction_site",
    "объект из перечня": "listed_object",
}

# flag code → (title, where the sub-column lives in the sheet)
FLAGS: dict[str, tuple[str, str]] = {
    "injured": ("Пострадавшие / погибшие", "ОДС ПСЦ, МВД, СМП, ЦЭМП, Мосгортранс"),
    "no_access": ("Нет доступа", "Служба 101, ОДС ПСЦ"),
    "threat": ("Угроза людям", "ОДС ПСЦ, ЦЭМП"),
    "offense": ("Правонарушение", "Классификатор МВД"),
    "not_on_site": ("Пострадавшие не на месте", "Классификатор СМП"),
    "gasification": ("Газификация", "Классификатор Мосгаз"),
    "medical_help": ("Требуется медицинская помощь", "ЦЭМП"),
    "evacuation": ("Требуется эвакуация", "ЦЭМП"),
    "mass_incident": ("Более 5 человек / особо опасно", "Классификатор ФСБ"),
    "road_closed": ("Перекрытие движения", "Мосгортранс, ГКУ Организатор перевозок"),
    "tunnel": ("Тоннель", "Гормост"),
    "pedestrian_bridge": ("Пешеходный мост", "Гормост"),
    "road_bridge": ("Автомобильный мост", "Гормост"),
    "telecom_object": ("На объектах связи", "МГТС"),
    "construction_site": ("Стройка", "Департамент строительства"),
    "listed_object": ("Объект из перечня", "Департамент культуры"),
}

# Column M values → service codes.
MAIN_SERVICE_CODES: dict[str, str] = {
    "mchs": "101",
    "police": "102",
    "ambulance": "103",
    "mosgaz": "mosgaz",
    "moslift": "moslift",
    "autoroads": "autoroads",
    "mosvodocanal": "mosvodokanal",
    "metro": "metro",
    "oek": "oek",
    "mosgortrans": "mosgortrans",
    "moesk": "moesk",
    "moek": "moek",
    "mzd": "mzd",
    "mgts": "mgts",
    "mosvodostok": "mosvodostok",
    "moscollector": "moscollector",
    "gormost": "gormost",
    "gkh": "gkh",
    "dep.tszn": "dep_tszn",
    "zodd": "codd",
    "msppn": "msppn",
    "depeco": "dep_ppoos",
    "zemp": "cemp",
    "мср": "msr_kub",
}

# Services whose cards reach them through an integrated information system, not АРМ-112
# (memo, page 7). Everyone else works on АРМ-112 and is the audience of the trainer.
INTEGRATED_SERVICES = {"101", "102", "103", "codd", "mosvodokanal", "dgp_integration"}
NO_REJECT_SERVICES = {"103"}


@dataclass
class ServiceColumn:
    index: int  # 1-based column index in the sheet
    letter: str
    service_code: str
    flag: str | None


@dataclass
class ServiceRule:
    service: str
    when: list[str]
    notify: bool
    value: str  # raw cell text, kept for traceability

    def as_dict(self) -> dict:
        return {
            "service": self.service,
            "when": self.when,
            "notify": self.notify,
            "value": self.value,
        }


@dataclass
class IncidentType:
    code: str
    group_code: str
    stat_group: str | None
    sign1: str
    sign2: str | None
    sign3: str | None
    hints: str | None
    final_title: str
    ekp_title: str | None
    main_service: str | None
    main_service_raw: str | None
    service_rules: list[ServiceRule]
    flag_codes: list[str]
    source_row: int

    @property
    def default_services(self) -> list[str]:
        return sorted({r.service for r in self.service_rules if r.notify and not r.when})


@dataclass
class IncidentGroup:
    code: str
    title: str
    number: int
    source_row: int


@dataclass
class Classifier:
    groups: list[IncidentGroup]
    types: list[IncidentType]
    columns: list[ServiceColumn]
    unparsed_columns: list[tuple[str, str, str, str]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def service_codes(self) -> list[str]:
        seen: dict[str, None] = {}
        for c in self.columns:
            seen.setdefault(c.service_code)
        return list(seen)

    def type_by_code(self) -> dict[str, IncidentType]:
        return {t.code: t for t in self.types}


def normalize(text: object) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()


def _key(text: object) -> str:
    return normalize(text).lower()


def service_catalog(codes: list[str]) -> list[dict]:
    """Rows for the ``services`` table in sheet order."""
    titles = {code: (title, short) for code, title, short in SERVICE_COLUMNS.values()}
    rows = []
    for code in codes:
        title, short = titles[code]
        rows.append(
            {
                "code": code,
                "title": title,
                "short_title": short,
                "no_reject": code in NO_REJECT_SERVICES,
                "via_arm112": code not in INTEGRATED_SERVICES,
            }
        )
    return rows


def _header_cells(ws) -> dict[tuple[int, int], str]:
    """Resolves merged header cells so every (row, col) has its text."""
    origin: dict[tuple[int, int], tuple[int, int]] = {}
    for merged in ws.merged_cells.ranges:
        if merged.min_row > HEADER_ROWS:
            continue
        for r in range(merged.min_row, merged.max_row + 1):
            for c in range(merged.min_col, merged.max_col + 1):
                origin[(r, c)] = (merged.min_row, merged.min_col)
    cells = {}
    for r in range(1, HEADER_ROWS + 1):
        for c in range(1, ws.max_column + 1):
            src = origin.get((r, c), (r, c))
            cells[(r, c)] = normalize(ws.cell(*src).value)
    return cells


def parse_service_columns(ws) -> tuple[list[ServiceColumn], list[tuple[str, str, str, str]]]:
    from openpyxl.utils import get_column_letter

    cells = _header_cells(ws)
    columns: list[ServiceColumn] = []
    unparsed: list[tuple[str, str, str, str]] = []
    for c in range(FIRST_SERVICE_COLUMN, ws.max_column + 1):
        family, sub, cond = cells[(1, c)], cells[(2, c)], cells[(3, c)]
        letter = get_column_letter(c)
        if not family:
            continue
        # A vertically merged header repeats the family text on lower rows.
        if _key(sub) == _key(family):
            sub = ""
        if _key(cond) == _key(sub):
            cond = ""
        service = SERVICE_COLUMNS.get((_key(family), _key(sub)))
        if service is None and sub:
            # Row 2 may carry the condition instead of a sub-service (МВД, СМП).
            service = SERVICE_COLUMNS.get((_key(family), ""))
            if service is not None and not cond:
                cond = sub
        if service is None or _key(cond) not in CONDITION_FLAGS:
            unparsed.append((letter, family, sub, cond))
            continue
        columns.append(ServiceColumn(c, letter, service[0], CONDITION_FLAGS[_key(cond)]))
    return columns, unparsed


def _cell_value(raw: object, final_title: str) -> str:
    text = normalize(raw)
    # «=K5»-style formulas simply repeat the final type of the row.
    if text.startswith("=K"):
        return final_title
    return text


def parse_classifier(path: Path) -> Classifier:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=False, data_only=False)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"в файле нет листа «{SHEET_NAME}»: {wb.sheetnames}")
    ws = wb[SHEET_NAME]
    columns, unparsed = parse_service_columns(ws)
    groups: list[IncidentGroup] = []
    types: list[IncidentType] = []
    warnings: list[str] = []
    current_group: IncidentGroup | None = None
    stat_group: str | None = None
    seen_codes: Counter[str] = Counter()

    for row_index, row in enumerate(
        ws.iter_rows(min_row=HEADER_ROWS + 1, values_only=True), start=HEADER_ROWS + 1
    ):
        final_title = normalize(row[COL_FINAL - 1])
        group_title = normalize(row[COL_STAT_GROUP - 1])
        if not final_title:
            if group_title and not normalize(row[COL_SIGN1 - 1]):
                number = row[COL_NUMBER - 1]
                # The last group («БПЛА») has no number in column E; numbering continues.
                code = int(number) if isinstance(number, int | float) else len(groups) + 1
                current_group = IncidentGroup(str(code), group_title, code, row_index)
                groups.append(current_group)
                stat_group = None
            elif any(v is not None for v in row):
                warnings.append(f"строка {row_index}: нет итогового типа, пропущена")
            continue
        if current_group is None:
            warnings.append(f"строка {row_index}: тип до первой группы, пропущена")
            continue
        codes = [row[i] for i in range(COL_G - 1, COL_P3)]
        if any(v is None for v in codes):
            warnings.append(f"строка {row_index}: нет кодов Г/п1/п2/п3, пропущена")
            continue
        code = ".".join(str(int(v)) for v in codes)
        seen_codes[code] += 1
        if seen_codes[code] > 1:
            code = f"{code}#{seen_codes[code]}"
            warnings.append(f"строка {row_index}: повтор кода, использован {code}")
        if group_title:
            stat_group = group_title
        rules: list[ServiceRule] = []
        flags: set[str] = set()
        for column in columns:
            value = _cell_value(row[column.index - 1], final_title)
            if not value:
                continue
            notify = value.lower() != NO_RESPONSE
            when = [column.flag] if column.flag else []
            rules.append(ServiceRule(column.service_code, when, notify, value))
            if column.flag:
                flags.add(column.flag)
        main_raw = normalize(row[COL_MAIN - 1]) or None
        main_codes = [
            MAIN_SERVICE_CODES.get(_key(part)) for part in (main_raw or "").split(",") if _key(part)
        ]
        if main_raw and any(c is None for c in main_codes):
            warnings.append(f"строка {row_index}: неизвестная главная служба «{main_raw}»")
        types.append(
            IncidentType(
                code=code,
                group_code=current_group.code,
                stat_group=stat_group,
                sign1=normalize(row[COL_SIGN1 - 1]),
                sign2=normalize(row[COL_SIGN2 - 1]) or None,
                sign3=normalize(row[COL_SIGN3 - 1]) or None,
                hints=normalize(row[COL_HINTS - 1]) or None,
                final_title=final_title,
                ekp_title=normalize(row[COL_EKP - 1]) or None,
                main_service=next((c for c in main_codes if c), None),
                main_service_raw=main_raw,
                service_rules=rules,
                flag_codes=sorted(flags),
                source_row=row_index,
            )
        )
    return Classifier(groups, types, columns, unparsed, warnings)
