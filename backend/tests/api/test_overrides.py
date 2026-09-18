"""Teacher's review actions (wave 9): the override of a total with a mandatory reason, the
comments, what the trainee sees, the events, the audit and the report with remarks."""

import uuid

import pytest
from httpx import AsyncClient
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncEngine

from app.db import SessionLocal
from app.events import events_after
from tests.api.conftest import DATA_DIR
from tests.api.test_attempts import journal, make_session, set_status
from tests.conftest import bearer, login

pytestmark = pytest.mark.skipif(
    not (DATA_DIR / "seed" / "classifier.json").exists(),
    reason="нет data/seed/classifier.json",
)

CHAIN = [
    ("accepted", {}),
    ("response_started", {"order_number": "14-217", "comment": "Направлен дежурный слесарь"}),
    ("arrived", {}),
    ("works_started", {"comment": "Мусоропровод вскрыт"}),
    ("works_done", {"comment": "Задымление устранено, ствол промыт, пострадавших нет"}),
]


async def evaluated_attempt(client: AsyncClient) -> tuple[str, str, dict]:
    """A closed and scored card of student1; returns (session_id, attempt_id, student token)."""
    session_id = str(await make_session(scenario_keys=["card_2-1_zadymlenie_musoroprovoda"]))
    student = await login(client, "student1")
    attempt_id = (await journal(client, student, uuid.UUID(session_id)))["items"][0]["attempt_id"]
    await client.post(f"/api/attempts/{attempt_id}/open", headers=bearer(student))
    for status, extra in CHAIN:
        r = await set_status(client, student, attempt_id, status=status, **extra)
        assert r.status_code == 200, r.text
    return session_id, attempt_id, student


async def _audit_actions(engine: AsyncEngine) -> list[str]:
    async with engine.connect() as conn:
        rows = await conn.execute(text("SELECT action FROM audit_log ORDER BY id"))
        return [r[0] for r in rows]


async def test_override_needs_reason_and_is_visible_to_student(
    client: AsyncClient, admin_engine: AsyncEngine
) -> None:
    session_id, attempt_id, student = await evaluated_attempt(client)
    teacher = await login(client, "teacher1")
    before = (await client.get(f"/api/attempts/{attempt_id}", headers=bearer(teacher))).json()
    old_total = before["evaluation"]["total"]
    assert before["override"] is None

    # Without a reason: rejected, nothing changes.
    r = await client.patch(
        f"/api/attempts/{attempt_id}/evaluation",
        headers=bearer(teacher),
        json={"new_total": 55, "reason": ""},
    )
    assert r.status_code == 422, r.text
    r = await client.patch(
        f"/api/attempts/{attempt_id}/evaluation",
        headers=bearer(teacher),
        json={"new_total": 55, "reason": "  "},
    )
    assert r.status_code == 422, r.text
    assert "evaluation.override" not in await _audit_actions(admin_engine)

    r = await client.patch(
        f"/api/attempts/{attempt_id}/evaluation",
        headers=bearer(teacher),
        json={"new_total": 55, "reason": "Комментарий к работам не по существу"},
    )
    assert r.status_code == 200, r.text
    override = r.json()
    assert override["old_total"] == old_total
    assert override["new_total"] == 55
    assert override["new_passed"] is False
    assert override["teacher_name"] == "Иванова Мария Петровна"

    # The trainee sees the new total, the old one for striking through and the reason.
    mine = (await client.get(f"/api/attempts/{attempt_id}", headers=bearer(student))).json()
    assert mine["evaluation"]["total"] == 55
    assert mine["evaluation"]["passed"] is False
    assert mine["override"]["old_total"] == old_total
    assert mine["override"]["reason"] == "Комментарий к работам не по существу"

    # Audit with the reason, and the event for the live screens.
    async with admin_engine.connect() as conn:
        row = (
            await conn.execute(
                text("SELECT details FROM audit_log WHERE action = 'evaluation.override'")
            )
        ).one()
    assert row[0]["reason"] == "Комментарий к работам не по существу"
    assert row[0]["old_total"] == old_total
    async with SessionLocal() as session:
        events = await events_after(session, uuid.UUID(session_id), 0, student_id=None)
    assert [e.type for e in events][-1] == "evaluation.overridden"
    assert events[-1].payload["total"] == 55

    # The report carries the changed total and the reason among the remarks.
    r = await client.get(f"/api/sessions/{session_id}/report", headers=bearer(teacher))
    assert r.status_code == 200, r.text
    rows = [a for s in r.json()["students"] for a in s["attempts"] if a["id"] == attempt_id]
    assert rows[0]["total"] == 55
    assert rows[0]["overridden"] is True
    assert any("Оценка изменена преподавателем" in x for x in rows[0]["remarks"])


async def test_override_rules(client: AsyncClient) -> None:
    session_id, attempt_id, student = await evaluated_attempt(client)
    body = {"new_total": 80, "reason": "Проверено вручную"}
    # The trainee and the administrator cannot change a total.
    r = await client.patch(
        f"/api/attempts/{attempt_id}/evaluation", headers=bearer(student), json=body
    )
    assert r.status_code == 403
    admin = await login(client, "admin")
    r = await client.patch(
        f"/api/attempts/{attempt_id}/evaluation", headers=bearer(admin), json=body
    )
    assert r.status_code == 403
    # A colleague does not even learn that the attempt exists.
    other = await login(client, "teacher2")
    r = await client.patch(
        f"/api/attempts/{attempt_id}/evaluation", headers=bearer(other), json=body
    )
    assert r.status_code == 404
    # A card still in work has no total to change.
    page = await journal(client, student, uuid.UUID(session_id))
    open_ids = [i["attempt_id"] for i in page["items"] if i["state"] != "evaluated"]
    if open_ids:
        teacher = await login(client, "teacher1")
        r = await client.patch(
            f"/api/attempts/{open_ids[0]}/evaluation", headers=bearer(teacher), json=body
        )
        assert r.status_code == 409
        assert r.json()["error"]["code"] == "not_evaluated"


async def test_comments_reach_the_student(client: AsyncClient, admin_engine: AsyncEngine) -> None:
    session_id, attempt_id, student = await evaluated_attempt(client)
    teacher = await login(client, "teacher1")
    r = await client.post(
        f"/api/attempts/{attempt_id}/comments",
        headers=bearer(teacher),
        json={"text": "Хорошая цепочка, но наряд указан не в том поле."},
    )
    assert r.status_code == 201, r.text
    assert r.json()["author_name"] == "Иванова Мария Петровна"
    r = await client.post(
        f"/api/attempts/{attempt_id}/comments", headers=bearer(teacher), json={"text": "   "}
    )
    assert r.status_code == 422
    r = await client.post(
        f"/api/attempts/{attempt_id}/comments", headers=bearer(student), json={"text": "x"}
    )
    assert r.status_code == 403

    mine = (await client.get(f"/api/attempts/{attempt_id}", headers=bearer(student))).json()
    assert [c["text"] for c in mine["comments"]] == [
        "Хорошая цепочка, но наряд указан не в том поле."
    ]
    assert "attempt.comment" in await _audit_actions(admin_engine)
    async with SessionLocal() as session:
        events = await events_after(session, uuid.UUID(session_id), 0, student_id=None)
    assert events[-1].type == "attempt.commented"

    progress = (await client.get("/api/me/progress", headers=bearer(student))).json()
    row = next(s for s in progress["sessions"] if s["session_id"] == session_id)
    assert row["comments"] == 1
    assert row["evaluated"] == 1
    assert progress["average"] is not None


async def test_progress_is_for_students_only(client: AsyncClient) -> None:
    teacher = await login(client, "teacher1")
    r = await client.get("/api/me/progress", headers=bearer(teacher))
    assert r.status_code == 403


async def test_exports_pdf_xlsx_csv(client: AsyncClient, admin_engine: AsyncEngine) -> None:
    from io import BytesIO

    from openpyxl import load_workbook
    from pypdf import PdfReader

    session_id, attempt_id, student = await evaluated_attempt(client)
    teacher = await login(client, "teacher1")
    await client.patch(
        f"/api/attempts/{attempt_id}/evaluation",
        headers=bearer(teacher),
        json={"new_total": 61, "reason": "Слесарь направлен без номера наряда"},
    )

    r = await client.get(f"/api/sessions/{session_id}/report.pdf", headers=bearer(teacher))
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF")
    # Paragraphs wrap inside table cells: compare on collapsed whitespace.
    text_ = " ".join(
        " ".join(p.extract_text() or "" for p in PdfReader(BytesIO(r.content)).pages).split()
    )
    assert "Отчёт о занятии" in text_
    assert "Слесарь направлен без номера наряда" in text_
    assert "Обучающийся" in text_

    r = await client.get(f"/api/sessions/{session_id}/report.xlsx", headers=bearer(teacher))
    assert r.status_code == 200, r.text
    book = load_workbook(BytesIO(r.content))
    assert book.sheetnames == ["Итоги", "Обучающиеся", "Попытки"]
    attempts = list(book["Попытки"].iter_rows(values_only=True))
    assert attempts[0][0] == "Обучающийся"
    scored = [row for row in attempts[1:] if row[6] == "61"]
    assert scored and scored[0][16] == "да"

    r = await client.get(f"/api/sessions/{session_id}/report.csv", headers=bearer(teacher))
    assert r.status_code == 200, r.text
    body = r.content.decode("utf-8-sig")
    assert body.splitlines()[0].startswith("Обучающийся;Логин;Карточка")
    assert "Слесарь направлен без номера наряда" in body

    r = await client.get(f"/api/sessions/{session_id}/report.docx", headers=bearer(teacher))
    assert r.status_code == 404
    r = await client.get(f"/api/sessions/{session_id}/report.pdf", headers=bearer(student))
    assert r.status_code == 403
    other = await login(client, "teacher2")
    r = await client.get(f"/api/sessions/{session_id}/report.pdf", headers=bearer(other))
    assert r.status_code == 403  # a colleague's session (plan wave 4)
    assert (await _audit_actions(admin_engine)).count("report.export") == 3
