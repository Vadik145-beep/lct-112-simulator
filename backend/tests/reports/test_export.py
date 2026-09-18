"""Export of the report without a database: PDF with Cyrillic, XLSX sheets, CSV for Excel."""

import uuid
from datetime import UTC, datetime
from io import BytesIO

from openpyxl import load_workbook
from pypdf import PdfReader

from app.reports import export
from app.training.teacher_schemas import (
    ReportAttempt,
    ReportErrorCount,
    ReportOut,
    ReportStudent,
    ReportSummary,
)

NOW = datetime(2026, 9, 22, 10, 0, tzinfo=UTC)
ERROR = ReportErrorCount(code="no_order_number", title="Нет номера наряда", count=1)


def _attempt(total: float | None, **kw) -> ReportAttempt:
    base = dict(
        id=uuid.uuid4(),
        card_number="38260311",
        scenario_title="Задымление мусоропровода",
        incident_title="Задымление в жилом доме",
        state="evaluated" if total is not None else "issued",
        card_status="finished",
        card_status_title="Завершена",
        issued_at=NOW,
        submitted_at=NOW if total is not None else None,
        total=total,
        passed=None if total is None else total >= 70,
        seconds=41.5 if total is not None else None,
        norm_seconds=30,
        deviation=11.5 if total is not None else None,
        decision_expected="accept",
        decision_actual="accept" if total is not None else None,
        decision_correct=True if total is not None else None,
        errors=["Нет номера наряда"] if total is not None else [],
        grammar_percent=87.5 if total is not None else None,
        remarks=["Нет номера наряда: наряд не указан при начале реагирования"] if total else [],
    )
    base.update(kw)
    return ReportAttempt(**base)


def _report() -> ReportOut:
    student = ReportStudent(
        student_id=uuid.uuid4(),
        full_name="Кузнецова Анна Сергеевна",
        login="student2",
        service_code="territorial_oiv",
        attempts=[
            _attempt(
                64,
                overridden=True,
                override_reason="Комментарий к работам не по существу",
                comments=["Наряд указывайте в поле «Номер наряда»"],
                remarks=[
                    "Нет номера наряда: наряд не указан",
                    "Оценка изменена преподавателем: 91.0 → 64.0 "
                    "(Комментарий к работам не по существу)",
                    "Комментарий: Наряд указывайте в поле «Номер наряда»",
                ],
            ),
            _attempt(None),
        ],
        attempts_total=2,
        evaluated=1,
        passed=0,
        average=64.0,
        average_seconds=41.5,
        average_deviation=11.5,
        wrong_decisions=0,
        typical_errors=[
            ReportErrorCount(code="no_order_number", title="Нет номера наряда", count=1)
        ],
        grammar_percent=87.5,
    )
    return ReportOut(
        session_id=uuid.uuid4(),
        title="Проверка экспорта <и & спецсимволов>",
        status="finished",
        started_at=NOW,
        finished_at=NOW,
        norm_seconds=30,
        pass_threshold=70,
        summary=ReportSummary(
            students=1,
            participated=1,
            evaluated=1,
            passed=0,
            average=64.0,
            average_seconds=41.5,
            typical_errors=[
                ReportErrorCount(code="no_order_number", title="Нет номера наряда", count=1)
            ],
        ),
        students=[student],
    )


def test_pdf_has_cyrillic_text_and_remarks() -> None:
    data = export.to_pdf(_report())
    assert data.startswith(b"%PDF")
    text = " ".join(
        " ".join(p.extract_text() or "" for p in PdfReader(BytesIO(data)).pages).split()
    )
    assert "Отчёт о занятии: Проверка экспорта <и & спецсимволов>" in text
    assert "Кузнецова Анна Сергеевна" in text
    assert "Комментарий к работам не по существу" in text
    assert "* оценка изменена преподавателем" in text


def test_xlsx_sheets_and_rows() -> None:
    book = load_workbook(BytesIO(export.to_xlsx(_report())))
    assert book.sheetnames == ["Итоги", "Обучающиеся", "Попытки"]
    rows = list(book["Попытки"].iter_rows(values_only=True))
    assert rows[0][:3] == ("Обучающийся", "Логин", "Карточка")
    assert rows[1][6] == "64" and rows[1][16] == "да"
    assert rows[2][6] == "выдана"
    students = list(book["Обучающиеся"].iter_rows(values_only=True))
    assert students[1][0] == "Кузнецова Анна Сергеевна"
    assert students[1][10] == "Нет номера наряда ×1"


def test_csv_is_excel_friendly() -> None:
    data = export.to_csv(_report())
    assert data.startswith("\ufeff".encode())
    lines = data.decode("utf-8-sig").splitlines()
    assert lines[0].split(";")[:2] == ["Обучающийся", "Логин"]
    assert lines[1].startswith("Кузнецова Анна Сергеевна;student2;38260311")
    assert "Оценка изменена преподавателем" in lines[1]
