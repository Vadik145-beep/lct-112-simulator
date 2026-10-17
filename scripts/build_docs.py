"""Сборка документации.

    python scripts/build_docs.py . docs/submission/Документация.docx       # общий комплект в DOCX
    python scripts/build_docs.py . docs/submission --pdf --only guides      # четыре PDF из docs/guides

PDF печатает Chromium из пакета playwright для Python, нужна версия 131 и новее (номера страниц
в полях страницы). Путь к другому Chromium: --browser или переменная DOCS_CHROMIUM. Шрифты
PT Sans и PT Mono лежат в docs/fonts, из интернета ничего не загружается.
"""

import argparse
import html
import os
import re
import tempfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

FILES = [
    ("Пояснительная записка", "docs/EXPLANATORY_NOTE.md"),
    ("Маршрут проверки решения", "docs/EXPERT_GUIDE.md"),
    ("Архитектура решения", "docs/ARCHITECTURE.md"),
    ("Методы обработки данных", "docs/METHODS.md"),
    ("Установка, настройка и восстановление", "docs/INSTALL.md"),
    ("Руководство пользователя", "docs/USER_GUIDE.md"),
    ("Данные заказчика", "docs/DATASET.md"),
    ("Условия и ограничения решения", "docs/LIMITATIONS.md"),
    ("Производительность и нагрузка", "docs/PERFORMANCE.md"),
    ("Проверка подсчёта оценки на живых звонках", "docs/VOICE_TESTS.md"),
    ("Использованные библиотеки и компоненты", "docs/LIBRARIES.md"),
    ("Соответствие техническому заданию", "docs/REQUIREMENTS_MATRIX.md"),
]

LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
BOLD = re.compile(r"\*\*([^*]+)\*\*")
CODE = re.compile(r"`([^`]+)`")


def add_runs(par, text):
    """Пишет текст с учётом **жирного** и `кода`."""
    text = LINK.sub(r"\1", text)
    pos = 0
    pattern = re.compile(r"\*\*([^*]+)\*\*|`([^`]+)`")
    for m in pattern.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos : m.start()])
        if m.group(1) is not None:
            par.add_run(m.group(1)).bold = True
        else:
            r = par.add_run(m.group(2))
            r.font.name = "Consolas"
            r.font.size = Pt(10)
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def add_toc(doc):
    par = doc.add_paragraph()
    run = par.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Оглавление обновится при открытии документа."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr, fld_sep, placeholder, fld_end):
        run._r.append(el)


def flush_paragraph(doc, buf):
    if not buf:
        return
    par = doc.add_paragraph()
    add_runs(par, " ".join(buf))
    buf.clear()


def add_table(doc, rows):
    rows = [r for r in rows if not re.match(r"^\|[\s:|-]+\|$", r.strip())]
    if not rows:
        return
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    width = max(len(r) for r in cells)
    table = doc.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    for n, row in enumerate(cells):
        cs = table.add_row().cells
        for j in range(width):
            text = row[j] if j < len(row) else ""
            par = cs[j].paragraphs[0]
            add_runs(par, text)
            for run in par.runs:
                run.font.size = Pt(9)
                if n == 0:
                    run.bold = True


def render(doc, md, root):
    lines = md.splitlines()
    buf, table_buf, code_buf = [], [], []
    in_code = False
    for line in lines:
        if line.strip().startswith("```"):
            if in_code:
                flush_paragraph(doc, buf)
                par = doc.add_paragraph()
                run = par.add_run("\n".join(code_buf))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                code_buf.clear()
            in_code = not in_code
            continue
        if in_code:
            code_buf.append(line)
            continue
        if line.strip().startswith("|"):
            flush_paragraph(doc, buf)
            table_buf.append(line)
            continue
        if table_buf:
            add_table(doc, table_buf)
            table_buf = []
        img = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", line)
        if img:
            flush_paragraph(doc, buf)
            path = (root / "docs" / img.group(2)).resolve()
            if path.exists():
                doc.add_picture(str(path), width=Cm(16))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                if img.group(1):
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = cap.add_run(img.group(1))
                    run.italic = True
                    run.font.size = Pt(10)
            continue
        m = re.match(r"^(#{2,6})\s+(.*)$", line)
        if m:
            flush_paragraph(doc, buf)
            level = min(len(m.group(1)), 4)
            doc.add_heading(LINK.sub(r"\1", BOLD.sub(r"\1", m.group(2))), level=level)
            continue
        if re.match(r"^\s*[-*]\s+", line):
            flush_paragraph(doc, buf)
            par = doc.add_paragraph(style="List Bullet")
            add_runs(par, re.sub(r"^\s*[-*]\s+", "", line))
            continue
        if re.match(r"^\s*\d+\.\s+", line):
            flush_paragraph(doc, buf)
            par = doc.add_paragraph(style="List Number")
            add_runs(par, re.sub(r"^\s*\d+\.\s+", "", line))
            continue
        if not line.strip():
            flush_paragraph(doc, buf)
            continue
        buf.append(line.strip())
    flush_paragraph(doc, buf)
    if table_buf:
        add_table(doc, table_buf)


def build_docx(root: Path, out: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    for section in doc.sections:
        section.left_margin = Cm(2)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)

    # титульный лист
    for _ in range(6):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Тренажёр оператора ДДС-112")
    run.bold = True
    run.font.size = Pt(24)
    for text, size, bold in [
        ("Сопроводительная документация к решению", 14, False),
        ("", 12, False),
        ("Хакатон «Лидеры цифровой трансформации 2026»", 12, False),
        (
            "Задача № 9. Учебное программное обеспечение для подготовки оператора ДДС города Москвы "
            "с использованием искусственного интеллекта",
            12,
            False,
        ),
        ("", 12, False),
        (
            "Постановщик: Департамент по делам гражданской обороны, чрезвычайным ситуациям "
            "и пожарной безопасности города Москвы, ГБУ «Система 112»",
            12,
            False,
        ),
        ("", 12, False),
        ("Команда «Ультратех»", 12, True),
        (date.today().strftime("%d.%m.%Y"), 12, False),
    ]:
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Содержание", level=1)
    add_toc(doc)

    for n, (title_text, rel) in enumerate(FILES, 1):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading(f"{n}. {title_text}", level=1)
        md = (root / rel).read_text(encoding="utf-8")
        md = re.sub(r"^#\s+.*$", "", md, count=1, flags=re.M)
        render(doc, md, root)

    doc.save(out)
    print("DOCX собран:", out)


# ---------- PDF: четыре отдельных документа из docs/guides ----------

GUIDES = [
    ("docs/guides/TECH_DESCRIPTION.md", "1_Общее_техническое_описание.pdf", "Общее техническое описание"),
    ("docs/guides/ADMIN_GUIDE.md", "2_Руководство_администратора.pdf", "Руководство администратора"),
    ("docs/guides/TEACHER_GUIDE.md", "3_Руководство_преподавателя.pdf", "Руководство преподавателя"),
    ("docs/guides/STUDENT_GUIDE.md", "4_Руководство_обучаемого.pdf", "Руководство обучаемого (диспетчер ДДС)"),
]
PDF_PREFIX = "Ультратех_задача9_"
PRODUCT = "Тренажёр диспетчера ДДС"
# Page margin boxes (the page number in the footer) appeared in Chromium 131.
MIN_CHROMIUM = 131
# A placeholder as wide as a real page number: the second pass must not move the pages.
PAGE_PLACEHOLDER = "000"

PDF_CSS = """
@font-face { font-family: "PT Sans"; src: url("{fonts}/PT_Sans-Web-Regular.ttf"); }
@font-face { font-family: "PT Sans"; font-weight: bold; src: url("{fonts}/PT_Sans-Web-Bold.ttf"); }
@font-face { font-family: "PT Sans"; font-style: italic; src: url("{fonts}/PT_Sans-Web-Italic.ttf"); }
@font-face { font-family: "PT Sans"; font-weight: bold; font-style: italic;
             src: url("{fonts}/PT_Sans-Web-BoldItalic.ttf"); }
@font-face { font-family: "PT Mono"; src: url("{fonts}/PTM55FT.ttf"); }
@page {
  size: A4; margin: 18mm 16mm 18mm 22mm;
  @top-right { content: "{running}"; font: 8pt "PT Sans"; color: #8a8a8a; }
  @bottom-center { content: counter(page); font: 9pt "PT Sans"; color: #555; }
}
@page :first { @top-right { content: none; } @bottom-center { content: none; } }
html { font-family: "PT Sans", sans-serif; font-size: 11pt; line-height: 1.42; color: #1a1a1a; }
body { margin: 0; }
.title { height: 255mm; display: flex; flex-direction: column; break-after: page; }
.title .team { font-size: 12pt; color: #444; }
.title .middle { margin-top: 70mm; }
.title .product { font-size: 15pt; color: #c2410c; font-weight: bold; letter-spacing: .02em; }
.title .doc { font-size: 28pt; font-weight: bold; line-height: 1.15; margin: 4mm 0 10mm; }
.title .about { font-size: 11.5pt; color: #333; max-width: 140mm; }
.title .bottom { margin-top: auto; font-size: 10.5pt; color: #444; }
.toc { break-after: page; }
.toc .head { font-size: 18pt; font-weight: bold; margin-bottom: 6mm; }
.toc a { display: flex; color: inherit; text-decoration: none; margin: 1.6mm 0; }
.toc a .t { flex: none; max-width: 150mm; }
.toc a .dots { flex: 1; border-bottom: 1px dotted #999; margin: 0 2mm 1.2mm; }
.toc a.l3 { padding-left: 7mm; font-size: 10.5pt; color: #333; }
.toc a.l2 { font-weight: bold; margin-top: 3mm; }
h2 { font-size: 16pt; margin: 9mm 0 3mm; color: #111; break-after: avoid; }
h2.first { margin-top: 0; }
h3 { font-size: 13pt; margin: 6mm 0 2mm; break-after: avoid; }
h4 { font-size: 11.5pt; margin: 4mm 0 1.5mm; break-after: avoid; }
p { margin: 0 0 2.6mm; }
p:has(+ ul), p:has(+ ol), p:has(+ table), p:has(+ pre) { break-after: avoid; }
ul, ol { margin: 0 0 2.6mm; padding-left: 7mm; }
li { margin: 0 0 1.2mm; }
li > ul, li > ol { margin-top: 1.2mm; }
figure { margin: 3mm 0 5mm; break-inside: avoid; text-align: center; }
figure img { max-width: 100%; max-height: 200mm; border: 1px solid #cfd4da; }
figcaption { font-size: 9.5pt; font-style: italic; color: #555; margin-top: 1.5mm; }
table { border-collapse: collapse; width: 100%; margin: 2mm 0 4mm; font-size: 9.5pt; }
th, td { border: 1px solid #c9ced4; padding: 1.4mm 2mm; text-align: left; vertical-align: top; }
th { background: #eef1f4; }
tr { break-inside: avoid; }
code { font-family: "PT Mono", monospace; font-size: 9.5pt; background: #f1f3f5; padding: 0 1mm; }
pre { font-family: "PT Mono", monospace; font-size: 9pt; background: #f1f3f5; padding: 2.5mm 3mm;
      white-space: pre-wrap; break-inside: avoid; border-left: 2px solid #c2410c; }
pre code { background: none; padding: 0; }
blockquote { margin: 0 0 3mm; padding: 2mm 4mm; background: #fff7ed; border-left: 3px solid #c2410c; }
a { color: #1d4ed8; }
"""

LIST_ITEM = re.compile(r"^(\s*)([-*]|\d+\.)\s+(.*)$")
IMAGE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
HEADING = re.compile(r"^(#{2,4})\s+(.*)$")


def inline_html(text: str) -> str:
    """**жирный**, `код` и ссылки в HTML; ссылки на файлы репозитория остаются текстом."""
    out = []
    for part in re.split(r"(`[^`]+`)", text):
        if len(part) > 1 and part.startswith("`") and part.endswith("`"):
            out.append(f"<code>{html.escape(part[1:-1])}</code>")
            continue
        part = html.escape(part, quote=False)
        part = LINK.sub(
            lambda m: f'<a href="{m.group(2)}">{m.group(1)}</a>'
            if m.group(2).startswith(("http://", "https://"))
            else m.group(1),
            part,
        )
        out.append(BOLD.sub(r"<b>\1</b>", part))
    return "".join(out)


def md_to_html(md: str, base: Path) -> tuple[str, list[tuple[int, str, str]]]:
    """Markdown руководства в HTML. Возвращает тело и заголовки (уровень, id, текст) для оглавления."""
    md = re.sub(r"<!--.*?-->", "", md, flags=re.S)
    md = re.sub(r"^#\s+.*$", "", md, count=1, flags=re.M)
    out: list[str] = []
    headings: list[tuple[int, str, str]] = []
    para: list[str] = []
    table: list[str] = []
    code: list[str] | None = None
    lists: list[tuple[int, str]] = []  # (indent, "ul" | "ol") of the open lists

    def flush_para():
        if para:
            text = inline_html(" ".join(para))
            out.append(f"<br>{text}" if lists else f"<p>{text}</p>")
            para.clear()

    def close_lists(indent: int = -1):
        while lists and lists[-1][0] > indent:
            out.append(f"</li></{lists.pop()[1]}>")

    def flush_table():
        if not table:
            return
        rows = [r for r in table if not re.match(r"^\|[\s:|-]+\|$", r.strip())]
        cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        out.append("<table><thead><tr>" + "".join(f"<th>{inline_html(c)}</th>" for c in cells[0]))
        out.append("</tr></thead><tbody>")
        for row in cells[1:]:
            out.append("<tr>" + "".join(f"<td>{inline_html(c)}</td>" for c in row) + "</tr>")
        out.append("</tbody></table>")
        table.clear()

    for line in md.splitlines():
        if code is not None:
            if line.strip().startswith("```"):
                out.append(f"<pre><code>{html.escape(chr(10).join(code))}</code></pre>")
                code = None
            else:
                code.append(line)
            continue
        if line.strip().startswith("```"):
            flush_para(), flush_table(), close_lists()
            code = []
            continue
        if line.strip().startswith("|"):
            flush_para(), close_lists()
            table.append(line)
            continue
        flush_table()
        if not line.strip():
            flush_para()
            continue
        m = HEADING.match(line)
        if m:
            flush_para(), close_lists()
            level = len(m.group(1))
            hid = f"h{len(headings) + 1}"
            headings.append((level, hid, BOLD.sub(r"\1", m.group(2)).strip()))
            cls = ' class="first"' if len(headings) == 1 else ""
            out.append(f'<h{level} id="{hid}"{cls}>{inline_html(m.group(2))}</h{level}>')
            continue
        m = IMAGE.match(line.strip())
        if m:
            flush_para(), close_lists()
            src = (base / m.group(2)).resolve()
            if not src.exists():
                raise SystemExit(f"нет картинки {src}")
            cap = f"<figcaption>{inline_html(m.group(1))}</figcaption>" if m.group(1) else ""
            out.append(f'<figure><img src="{src.as_uri()}">{cap}</figure>')
            continue
        m = LIST_ITEM.match(line)
        if m:
            flush_para()
            indent, kind = len(m.group(1)), "ol" if m.group(2)[0].isdigit() else "ul"
            close_lists(indent)
            if lists and lists[-1][0] == indent and lists[-1][1] != kind:
                close_lists(indent - 1)
            if lists and lists[-1][0] == indent:
                out.append(f"</li><li>{inline_html(m.group(3))}")
            else:
                start = f' start="{m.group(2)[:-1]}"' if kind == "ol" and m.group(2) != "1." else ""
                lists.append((indent, kind))
                out.append(f"<{kind}{start}><li>{inline_html(m.group(3))}")
            continue
        if line.startswith(">"):
            flush_para(), close_lists()
            out.append(f"<blockquote>{inline_html(line.lstrip('> '))}</blockquote>")
            continue
        if lists and not line.startswith((" ", "\t")):
            flush_para(), close_lists()
        if lists and not para:
            # A continuation line of a list item goes on in the same item.
            out.append(" " + inline_html(line.strip()))
            continue
        para.append(line.strip())
    flush_para(), flush_table(), close_lists()
    return "\n".join(out), headings


def guide_html(title: str, body: str, headings, pages: dict[str, int], fonts: Path) -> str:
    toc = []
    for level, hid, text in headings:
        if level > 3:
            continue
        toc.append(
            f'<a class="l{level}" href="#{hid}"><span class="t">{html.escape(text)}</span>'
            f'<span class="dots"></span><span>{pages.get(hid, PAGE_PLACEHOLDER)}</span></a>'
        )
    css = PDF_CSS.replace("{fonts}", fonts.as_uri()).replace("{running}", f"{PRODUCT}. {title}")
    today = date.today().strftime("%d.%m.%Y")
    return f"""<!doctype html><html lang="ru"><head><meta charset="utf-8"><style>{css}</style></head>
<body>
<div class="title">
  <div class="team">Команда «Ультратех», задача № 9</div>
  <div class="middle">
    <div class="product">{PRODUCT}</div>
    <div class="doc">{html.escape(title)}</div>
    <div class="about">Учебное программное обеспечение для подготовки диспетчеров
    дежурно-диспетчерских служб города Москвы с использованием искусственного интеллекта.</div>
  </div>
  <div class="bottom">Хакатон «Лидеры цифровой трансформации 2026»<br>
  Москва, редакция от {today}</div>
</div>
<div class="toc"><div class="head">Содержание</div>{"".join(toc)}</div>
{body}
</body></html>"""


def find_chromium() -> str | None:
    """Свежий Chromium из каталога браузеров playwright: встроенный в пакет бывает старше 131."""
    if os.environ.get("DOCS_CHROMIUM"):
        return os.environ["DOCS_CHROMIUM"]
    dirs = [Path.home() / ".cache" / "ms-playwright", Path.home() / "Library" / "Caches" / "ms-playwright"]
    if os.environ.get("PLAYWRIGHT_BROWSERS_PATH"):
        dirs.append(Path(os.environ["PLAYWRIGHT_BROWSERS_PATH"]))
    if os.environ.get("LOCALAPPDATA"):
        dirs.append(Path(os.environ["LOCALAPPDATA"]) / "ms-playwright")
    executables = ("chrome-win64/chrome.exe", "chrome-win/chrome.exe", "chrome-linux64/chrome",
                   "chrome-linux/chrome", "chrome-mac/Chromium.app/Contents/MacOS/Chromium")
    found = [
        (int(c.name.split("-")[-1]), str(c / exe))
        for d in dirs
        for c in d.glob("chromium-*")
        if c.name.split("-")[-1].isdigit()
        for exe in executables
        if (c / exe).exists()
    ]
    return max(found)[1] if found else None


def outline_pages(pdf: Path) -> list[int]:
    """Страницы закладок PDF по порядку (с единицы, титул — первая страница)."""
    from pypdf import PdfReader

    reader = PdfReader(pdf)
    pages: list[int] = []

    def walk(items):
        for item in items:
            if isinstance(item, list):
                walk(item)
            else:
                pages.append(reader.get_destination_page_number(item) + 1)

    walk(reader.outline)
    return pages


def build_guides_pdf(root: Path, out_dir: Path, browser_path: str | None) -> None:
    from playwright.sync_api import sync_playwright

    fonts = (root / "docs" / "fonts").resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as p, tempfile.TemporaryDirectory() as tmp:
        browser = p.chromium.launch(executable_path=browser_path or find_chromium())
        if int(browser.version.split(".")[0]) < MIN_CHROMIUM:
            raise SystemExit(
                f"Chromium {browser.version} не ставит номера страниц, нужен {MIN_CHROMIUM}+: "
                "python -m playwright install chromium или --browser <путь>"
            )
        page = browser.new_page()
        tmp_html = Path(tmp) / "guide.html"
        for rel, name, title in GUIDES:
            src = root / rel
            if not src.exists():
                print("  пропущен, нет файла:", rel)
                continue
            body, headings = md_to_html(src.read_text(encoding="utf-8"), src.parent)
            target = out_dir / f"{PDF_PREFIX}{name}"
            pages: dict[str, int] = {}
            # Two passes: the first finds the page of every heading from the PDF bookmarks,
            # the second prints the table of contents with those numbers.
            for _ in range(2):
                tmp_html.write_text(guide_html(title, body, headings, pages, fonts), encoding="utf-8")
                page.goto(tmp_html.as_uri())
                page.wait_for_load_state("networkidle")
                page.evaluate("document.fonts.ready")
                page.pdf(path=target, prefer_css_page_size=True, print_background=True,
                         outline=True, tagged=True)
                found = outline_pages(target)
                if len(found) != len(headings):
                    raise SystemExit(f"{rel}: закладок {len(found)}, заголовков {len(headings)}")
                pages = {hid: n for (_, hid, _), n in zip(headings, found)}
            print("  PDF:", target)
        browser.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Сборка документации в DOCX или PDF")
    parser.add_argument("root", type=Path, help="корень репозитория")
    parser.add_argument("out", type=Path, help="файл DOCX или папка для PDF")
    parser.add_argument("--pdf", action="store_true", help="собрать PDF вместо DOCX")
    parser.add_argument("--only", choices=["guides"], default="guides", help="комплект PDF")
    parser.add_argument("--browser", help="путь к Chromium 131+")
    args = parser.parse_args()
    if args.pdf:
        build_guides_pdf(args.root.resolve(), args.out, args.browser)
    else:
        build_docx(args.root, args.out)


if __name__ == "__main__":
    main()
